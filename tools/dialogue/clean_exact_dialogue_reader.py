#!/usr/bin/env python3
"""Clean the generated Chapters 0–3 reader without changing spoken dialogue.

The reader is a story-facing read-through, not a production document. This cleanup
removes writer/gameplay/implementation prose while preserving every actual spoken line
in exact order and wording.

Dialogue invariant: spoken lines use an ALL-CAPS speaker label in the first bold run,
ending in a colon. Title-case bold labels are production metadata, not dialogue.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
EXPECTED_DIALOGUE_LINES = 2901
LABEL_RE = re.compile(r"^.+:\s*$")

FULL_SKIP_HEADING_TERMS = (
    "production draft",
    "dialogue engine production draft",
    "person-brain",
    "role-balance note",
    "performance note",
    "performance check",
    "audit note",
    "audit check",
    "knowledge check",
    "knowledge checkpoint",
    "world-state checkpoint",
    "continuity note",
    "current production state",
    "conflict order",
    "gameplay",
    "optional gameplay",
    "boss handoff",
    "presentation check",
    "character check",
    "reveal-boundary check",
    "writer-facing",
    "implementation",
)

HEADING_EXACT = {
    "Noncombat Reset — Triage Line": "Triage Line",
    "Mid-Battle Transition — Tears Free": "Tears Free",
    "Boss Victory": "After the Battle",
    "Boss Combat — Briarhide Stalker": "Briarhide Stalker",
    "Optional Content — Available From Camp": "Camp",
    "Boss Battle — First Command Warden": "First Command Warden",
    "Optional Character-Life — Cresthaven Supply Room": "Cresthaven Supply Room",
    "Mandatory Interaction — Rest / Advance": "Rest / Advance",
    "Party Join": "Nimera Joins",
    "Return To Cleanup": "Back to Cresthaven",
    "Post-Boss": "After the Battle",
    "Next Morning — Player Preparation Window": "Next Morning",
}

HEADING_REMOVE = {
    "P04 end state",
    "P05 end state",
    "P06 end state",
    "P07 end state",
    "System / Party State",
    "Beat-1 end state",
    "Beat End",
    "Scene seed",
    "Boss Combat",
    "Boss Combat — Integrated State",
    "Boss Combat — Mobile State",
    "Battle Character",
    "Encounter-direction priorities",
}

PROSE_REWRITES = {
    "Current atomic dialogue edition. Spoken lines are copied verbatim from active production sources; writer-facing audits, implementation notes, and future-facing production metadata are omitted.":
        "Current atomic dialogue edition. Spoken lines are copied verbatim from the active source scenes; production notes are omitted.",
    "After the encounter, the convoy officer reaches Cyanis's side of the wreck field.":
        "After the fighting eases, the convoy officer reaches Cyanis's side of the wreck field.",
    "The Card remains inert for the entire encounter.": "The Card remains inert.",
    "Riftmaw and the War-Sorcerer are defeated in the same encounter.":
        "Riftmaw and the War-Sorcerer fall together.",
    "The final boss encounter is over.": "The fighting is over.",
    "The fight is lethal. The party defeats and kills the Briarhide Stalker through ordinary combat victory.":
        "The fight is lethal. The party kills the Briarhide Stalker.",
    "A wounded escort calls Cyanis over. The relevant NPC can appear for the scene.":
        "A wounded escort calls Cyanis over.",
    "The scene releases briefly into normal camp work.": "Camp work resumes briefly.",
    "The scene returns to the defended treatment area before P06.":
        "They return to the defended treatment area.",
    "Nobody proposes another separation attempt here. Chapter 1 owns the first deliberate transfer test.":
        "Nobody proposes another separation attempt.",
    "In the next visual beat, both Cyanis and Ilyra are shown passed out where they were sitting.":
        "A moment later, Cyanis and Ilyra are both asleep where they were sitting.",
    "Ilyra gives him the relevant version, not the whole convoy story.":
        "Ilyra gives him the relevant version, not the whole convoy account.",
    "Maevra joins the intake scene.": "Maevra joins them at intake.",
    "Greenhollow comes into view as part of normal field progression.": "Greenhollow comes into view.",
    "Torren knows these were people from a fort near his town. He does not become eloquent because the scene is sad.":
        "Torren knows these were people from a fort near his town.",
    "Torren studies the transition.": "Torren studies the join between the structures.",
    "Dunmere is functioning. Shops are open, people are working, and ordinary town life continues. The first thing the group sees is not a disaster scene.":
        "Dunmere is functioning. Shops are open, people are working, and ordinary town life continues. The group does not arrive to a disaster.",
    "The creature is defeated and no longer contests the chamber. Exact defeat animation/final physical state belongs to encounter presentation, but the route is no longer blocked by its active presence.":
        "The creature is defeated and no longer contests the chamber. The route is open.",
    "Nimera enters combat without a recruitment discussion or order.":
        "Nimera joins the fight without discussion.",
    "Ilyra examines the injury without turning the scene into a full treatment sequence.":
        "Ilyra examines the injury briefly.",
    "Rhazek retrieves himself enough to move. He remains visibly affected by the fight; the scene does not erase the defeat because he can still stand.":
        "Rhazek retrieves himself enough to move. He remains visibly affected by the fight.",
    "The controlled arrival stretch ends at the palace / royal-district transition.":
        "They reach the palace and royal district.",
    "The scene ends on controlled investigation rather than accusation.":
        "The room settles into controlled investigation rather than accusation.",
    "The keeper indicates the approved descent route. No elaborate map or prop animation is needed.":
        "The keeper indicates the approved descent route.",
    "During the investigation, Cyanis checks the Card once because Beat 11's change is still fresh.":
        "During the investigation, Cyanis checks the Card once; the change is still fresh.",
    "He brings out / retrieves the existing copy made from the Chapter-2 Ancient mural / route evidence.":
        "He brings out the existing copy of the Ancient mural and route evidence.",
    "A short departure trigger gathers the permanent four.": "The permanent four gather to depart.",
    "No new Card event occurs. Cyanis's Card remains stable deep Ruby and put away.":
        "Cyanis's Card remains stable deep Ruby and put away.",
    "BATTLE — ARCHIVE SCRIBE ENGINE": "The Archive Scribe Engine activates.",
}

DROP_EXACT = {
    "Cyanis and Ilyra are now the active combat pair.",
    "There is a real noncombat interval here.",
    "No immediate boss sting. No continuous Card ward. No second flare yet.",
    "Maevra's splint remains part of her visible Chapter-1 state.",
    "Available during the Chapter-3 cleanup window after Beat 15.",
    "Scene ends.",
    "The story trigger is short.",
    "Transition directly into:",
    "Current story presentation:",
    "Do not activate it yet.",
    "Do not inspect the room yet.",
    "Do not explain:",
    "They do not become default residents or traveling companions.",
    "No concealed Ruin Vanguard appears here.",
    "No Riftmaw appears here.",
    "No Card flare occurs here.",
    "No added transformation or hidden meaning.",
    "No Prime, Might, Last Sentinel, bearer, ancient-weapon, or Entity explanation is revealed.",
    "No incomplete protection, green-and-gold geometry, Card reaction, identity reveal, or transformation occurs.",
    "This is not a defeat-to-death sequence and not a chase setup.",
    "No plot reveal occurs here.",
    "This is the first time Ilyra is allowed to exist for several uninterrupted minutes without an active patient, battle, or Card response demanding her professional voice.",
    "No elaborate prop routine is needed.",
    "No confession follows. No relationship summary. No discussion of destiny, the Card, or what the day meant.",
    "The conversation leaves Torren's personal history alone. What interests Ilyra here is the absurdity of the argument, not diagnosing what Maevra feels about him.",
    "The exchange ends because the subject has run its natural course, not because every speaker has delivered one short line.",
    "The exchange establishes their familiarity without pausing to explain its history.",
    "No additional mechanism opens. No treasure-door or machinery sequence is added.",
    "The exchange ends because the group has finished amusing itself, not because Torren has reached a terseness quota.",
    "The protected old slut / old cut misunderstanding remains, but it functions as an early accidental jump into Cyanis and Torren's sharper humor rather than proof that their later brother-like relationship is already fully developed.",
    "This is not another ancient-Junction-map explanation and does not advance the Cistern Hunt.",
    "The treatment task is the circumstance that puts them alone together, not Ilyra's entire social function. Once the work is underway, Ilyra is allowed ordinary curiosity and gossip rather than therapist behavior.",
    "No Prime, Last Sentinel, bearer, Entity, or complete ancient-network explanation is added at camp.",
    "Nobody calls attention to Torren's pouch, weed, smoking supplies, or the later C06 setup.",
    "Optional short battle barks should stay practical and rare.",
    "The original Card-transport authorization and attached routing / expenditure records are placed before the room. Keep the physical presentation ordinary and administrative; the documents should look like real Crown business, not villain props.",
    "The assessment must visibly include Cyanis, Ilyra, Torren, and Nimera.",
    "No modern UI, targeting reticle, digital readout, or spoken identity check appears.",
    "Its battle identity should feel heavy, deliberate, chamber-bound, and built around physical magical wardcraft rather than modern machinery.",
    "No investigation state changes. No Card event occurs. No lore is advanced.",
    "There is no support wave and no injured Iron Cohort Soldier.",
}

PRODUCTION_PATTERNS = tuple(
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\b(?:P0?\d|C0?\d)\b",
        r"\bBeat\s+\d+\b",
        r"\bChapter[- ]\d+\b",
        r"\bthe chapter\b",
        r"\bplayer\b",
        r"\b(?:production|implementation|gameplay|writer-facing|authored|scripted)\b",
        r"\b(?:scene|dialogue|story)\b",
        r"\bCharacter-Life\b",
        r"\bHunt progression\b",
        r"\brecruitment\b",
        r"\bagency\b",
        r"\bsafety veto\b",
        r"\bcurrent atomic\b",
        r"\blegacy source key\b",
        r"\b(?:combat|battle) party\b",
        r"\bactive combat pair\b",
        r"\bpermanent combat(?:-capable)?\b",
        r"\bpermanent party (?:member|members|characters)\b",
        r"\bboss encounter\b",
        r"\bfinal boss\b",
        r"\bnormal boss combat\b",
        r"\bencounter-authorized\b",
        r"\bencounter authority\b",
        r"\bhard (?:scene|story) rule\b",
        r"\bplayable\b",
        r"\bcleanup window\b",
        r"\b(?:story|scene|departure) trigger\b",
        r"\btrigger fires\b",
        r"\bhandoff\b",
        r"\bobjective\b",
        r"\bpipeline\b",
        r"\bcontrol (?:returns|resumes|begins|pauses)\b",
        r"\bvisible traversal model\b",
        r"\bsparse practical dialogue is legal\b",
        r"\btimed dialogue\b",
        r"\bexchange (?:ends|establishes)\b",
        r"\bconversation leaves\b",
        r"\bdiagnosing\b",
        r"\bHP (?:bar|floor)\b",
        r"\bhealth bar\b",
        r"\bprotected HP\b",
        r"\bdisengagement state\b",
        r"\bforced nonlethal\b",
        r"\bnonlethal solution\b",
        r"\bdoes not participate in battle\b",
        r"\bdoes not transform into a corrupted form\b",
        r"\bNo ambiguity is introduced\b",
        r"\bstate transition\b",
        r"\bexact (?:mechanics|stats|attack tables|defeat animation|encounter-validation values)\b",
        r"\bfield models?\b",
        r"\bportraits?\b",
        r"\bdialogue (?:ui|layer|box)\b",
        r"\bcutscene\b",
        r"\bchoreography\b",
        r"\brandom[- ]encounter\b",
        r"\btutorial\b",
        r"\b(?:battle|encounter|field/story) presentation\b",
        r"\bcommandable party\b",
        r"\bnon-playable\b",
        r"\broad dungeon\b",
        r"\btravel-dialogue\b",
        r"\broute payoff\b",
        r"\bexploration-first\b",
        r"\bone-word quota\b",
        r"\bnot because the story\b",
        r"\bautomatic control of the decision\b",
        r"\bdoes not seize control\b",
        r"\bnot because Ilyra has taken control\b",
        r"\bretired records beat\b",
        r"\bdefault residents or traveling companions\b",
        r"\brequired (?:visual|traversal|combat)\b",
        r"\b(?:is|are) required\b",
        r"\bno emotional farewell\b",
        r"\breusable smoke loop\b",
        r"\bhand animation\b",
        r"\bpreparation window\b",
        r"\bplot reveal\b",
        r"\bprop routine\b",
        r"\brelationship summary\b",
        r"\bconfession follows\b",
        r"\bassessment must visibly include\b",
        r"\bmodern UI\b",
        r"\bbattle identity should feel\b",
        r"\binvestigation state changes\b",
        r"\blore is advanced\b",
        r"\btreasure-door\b",
        r"\ballowed to exist\b",
        r"\ballowed ordinary\b",
        r"\bmay appear again\b",
        r"^\s*END\s+C\d+\b",
        r"^\s*Do not\b",
        r"^\s*Use one\b",
        r"^\s*Short dialogue may\b",
        r"^\s*BOSS COMBAT BEGINS\b",
        r"^\s*Boss encounter begins\b",
        r"^\s*Combat continues\b",
        r"^\s*The battle ends\.?$",
        r"^\s*No new combat begins\.?$",
        r"^\s*Transition directly into:?$",
    )
)

RESIDUE_GUARD_PATTERNS = tuple(
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\b(?:P0?\d|C0?\d)\b",
        r"\bBeat\s+\d+\b",
        r"\bChapter[- ]\d+\b",
        r"\bthe chapter\b",
        r"\bplayer\b",
        r"\b(?:production|implementation|gameplay|writer-facing|authored|scripted)\b",
        r"\b(?:scene|dialogue|story)\b",
        r"\bCharacter-Life\b",
        r"\bHunt progression\b",
        r"\brecruitment\b",
        r"\bagency\b",
        r"\b(?:combat|battle) party\b",
        r"\bactive combat pair\b",
        r"\bpermanent party (?:member|members|characters)\b",
        r"\bboss encounter\b",
        r"\bencounter authority\b",
        r"\bhard (?:scene|story) rule\b",
        r"\bplayable\b",
        r"\bcleanup window\b",
        r"\bhandoff\b",
        r"\bobjective\b",
        r"\bHP (?:bar|floor)\b",
        r"\bhealth bar\b",
        r"\bfield models?\b",
        r"\bdialogue (?:ui|layer|box)\b",
        r"\brequired (?:visual|traversal|combat)\b",
        r"\bpreparation window\b",
        r"\bplot reveal\b",
        r"\bprop routine\b",
        r"\brelationship summary\b",
        r"\bassessment must visibly include\b",
        r"\bmodern UI\b",
        r"\bbattle identity should feel\b",
        r"\binvestigation state changes\b",
        r"\blore is advanced\b",
        r"^\s*END\s+C\d+\b",
        r"^\s*Do not\b",
        r"^\s*Use one\b",
    )
)


def is_heading(paragraph) -> bool:
    return bool(paragraph.style and paragraph.style.name and paragraph.style.name.startswith("Heading"))


def label_text(paragraph) -> str | None:
    if not paragraph.runs:
        return None
    first = paragraph.runs[0]
    if not first.bold or not LABEL_RE.match(first.text):
        return None
    return first.text.strip()[:-1].strip()


def is_dialogue(paragraph) -> bool:
    label = label_text(paragraph)
    if label is None:
        return False
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def is_non_dialogue_label(paragraph) -> bool:
    return label_text(paragraph) is not None and not is_dialogue(paragraph)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def clean_heading(text: str) -> str | None:
    text = text.strip()
    low = text.lower()
    if text in HEADING_REMOVE or re.fullmatch(r"c\d+\s+function", low):
        return None
    if low == "working royal audience":
        return "Royal Audience"
    if text in HEADING_EXACT:
        return HEADING_EXACT[text]

    text = re.sub(r"\s*[—-]\s*Scene$", "", text, flags=re.IGNORECASE).strip()

    for pattern in (
        r"^Major Story Interaction\s*[—-]\s*",
        r"^(?:First|Second|Third)?\s*Story Stop\s*[—-]\s*",
        r"^Short Story Stop\s*[—-]\s*",
        r"^Story Continuation\s*[—-]\s*",
        r"^Direct Cut\s*[—-]\s*",
        r"^Travel Cut\s*[—-]\s*",
        r"^Continuation\s*[—-]\s*",
        r"^Beat End\s*[—-]\s*",
        r"^Story Trigger\s*[—-]\s*",
        r"\s*[—-]\s*Story Trigger$",
        r"\s*[—-]\s*Optional Character-Life Trigger$",
        r"^Optional Character-Life Trigger\s*[—-]\s*",
        r"^Authored Stop\s*[—-]\s*",
        r"^Scripted Battle Event\s*[—-]\s*",
        r"^Authored(?:\s+Combat|\s+Opening|\s+Disengagement)?\s*[—-]?\s*",
        r"^Threshold Scene\s*[—-]\s*",
    ):
        changed = re.sub(pattern, "", text, flags=re.IGNORECASE).strip(" —-")
        if changed != text:
            text = changed or "Camp"

    low = text.lower()
    if low in {"beat end", "scene seed"}:
        return None
    if any(term in low for term in FULL_SKIP_HEADING_TERMS) or "guided traversal dialogue" in low:
        return None
    if "handoff" in low or low.startswith("pipeline:"):
        return None
    return text


def should_drop_prose(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if stripped in DROP_EXACT:
        return True
    return any(pattern.search(stripped) for pattern in PRODUCTION_PATTERNS)


def residual_meta(document) -> list[str]:
    found: list[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        if index <= 2 or is_heading(paragraph):
            continue
        text = paragraph.text.strip()
        if not text or is_dialogue(paragraph):
            continue
        if is_non_dialogue_label(paragraph) or any(pattern.search(text) for pattern in RESIDUE_GUARD_PATTERNS):
            found.append(text)
    return found


def main() -> int:
    if not READER.exists():
        raise RuntimeError(f"Generated reader does not exist: {READER.relative_to(ROOT)}")

    document = Document(READER)
    dialogue_before = [p.text for p in document.paragraphs if is_dialogue(p)]
    if len(dialogue_before) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(
            f"Generated reader contains {len(dialogue_before)} spoken dialogue lines; "
            f"expected {EXPECTED_DIALOGUE_LINES}. Refusing cleanup."
        )

    skip_non_dialogue = False
    for index, paragraph in enumerate(list(document.paragraphs)):
        text = paragraph.text.strip()

        if index <= 2:
            if text in PROSE_REWRITES:
                replace_paragraph_text(paragraph, PROSE_REWRITES[text])
            continue

        if is_heading(paragraph):
            cleaned = clean_heading(text)
            low = text.lower()
            if cleaned is None:
                remove_paragraph(paragraph)
                skip_non_dialogue = any(term in low for term in FULL_SKIP_HEADING_TERMS)
                continue
            skip_non_dialogue = False
            if cleaned != text:
                replace_paragraph_text(paragraph, cleaned)
            continue

        if is_dialogue(paragraph):
            continue

        if is_non_dialogue_label(paragraph):
            remove_paragraph(paragraph)
            continue

        if text in PROSE_REWRITES:
            replace_paragraph_text(paragraph, PROSE_REWRITES[text])
            continue

        if skip_non_dialogue or should_drop_prose(text):
            remove_paragraph(paragraph)

    dialogue_after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if dialogue_before != dialogue_after or len(dialogue_after) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(
            "Reader cleanup changed spoken dialogue; refusing to save "
            f"({len(dialogue_before)} before vs {len(dialogue_after)} after)."
        )

    residue = residual_meta(document)
    if residue:
        preview = "\n".join(f"- {line}" for line in residue[:12])
        raise RuntimeError(
            f"Reader cleanup left {len(residue)} unmistakable production-facing paragraphs:\n{preview}"
        )

    document.save(READER)
    print(
        f"Reader presentation cleanup complete; preserved {len(dialogue_after)} actual spoken dialogue lines "
        "exactly and passed the production-residue guard."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

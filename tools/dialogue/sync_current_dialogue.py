#!/usr/bin/env python3
"""Synchronize DIYSE Chapters 0-3 derived dialogue artifacts.

Standalone Markdown scene files under docs/03_DIALOGUE/PRODUCTION/CHAPTER_## are
exact wording authority. Current character/runtime/spec authorities must already be
source-closed before this tool is run.

This tool deliberately writes *derived artifacts only*:
- CHAPTER_##_REHEARSAL_FIRST_WORKING_DIALOGUE_MANUSCRIPT.md for Chapters 0-3
- docs/03_DIALOGUE/PRODUCTION/CHAPTER_0_3_DIALOGUE_SYNC_MANIFEST.md
- build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx

It does NOT modify chapter authority indexes or the Dialogue Master Index. Those
source-authority/status files must be promoted manually only after generated-output
verification succeeds.

Modes:
- --source-check : verify source closure, source selection, protected anchors, and
                   retired-canon guards without requiring derived outputs to exist.
- --check        : verify the generated Markdown manuscripts/manifest without writing.
- default        : regenerate Markdown/manifest and, unless --no-docx, the reader DOCX.

A generated DOCX still requires the normal render-and-visual-QA gate before delivery.
"""
from __future__ import annotations

import argparse
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
PROD = ROOT / "docs/03_DIALOGUE/PRODUCTION"
SYNC_MANIFEST = PROD / "CHAPTER_0_3_DIALOGUE_SYNC_MANIFEST.md"
SOURCE_CLOSURE = PROD / "CHAPTER_00_03_FULL_SOURCE_CLOSURE_2026-09-13.md"
READER_OUT = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"


@dataclass(frozen=True)
class SourceSpec:
    label: str
    pattern: str


@dataclass(frozen=True)
class ChapterSpec:
    chapter: str
    title: str
    sources: tuple[SourceSpec, ...]

    @property
    def directory(self) -> Path:
        return PROD / f"CHAPTER_{self.chapter}"

    @property
    def output(self) -> Path:
        return self.directory / f"CHAPTER_{self.chapter}_REHEARSAL_FIRST_WORKING_DIALOGUE_MANUSCRIPT.md"

    @property
    def index(self) -> Path:
        return self.directory / f"CHAPTER_{self.chapter}_DIALOGUE_AUTHORITY_INDEX.md"




CHAPTERS: tuple[ChapterSpec, ...] = (
    ChapterSpec(
        "00",
        "Chapter 0",
        tuple(SourceSpec(f"Beat {n}", f"CH00_B{n:02d}_*_DIALOGUE.md") for n in range(1, 8))
        + (SourceSpec("C01 — Six Minutes", "C01_SIX_MINUTES_DIALOGUE.md"),),
    ),
    ChapterSpec(
        "01",
        "Chapter 1",
        tuple(SourceSpec(f"Beat {n}", f"CH01_B{n:02d}_*_DIALOGUE.md") for n in range(1, 13))
        + (
            SourceSpec("C02 — Torren's Version of Dinner", "C02_TORRENS_VERSION_OF_DINNER_DIALOGUE.md"),
            SourceSpec("C03 — What the Map Says", "C03_WHAT_THE_MAP_SAYS_DIALOGUE.md"),
            SourceSpec("C04 — Not Professionally", "C04_NOT_PROFESSIONALLY_DIALOGUE.md"),
        ),
    ),
    ChapterSpec(
        "02",
        "Chapter 2",
        tuple(SourceSpec(f"Beat {n}", f"CH02_B{n:02d}_*_DIALOGUE.md") for n in range(1, 16))
        + (SourceSpec("C05 — Still Burns", "C05_STILL_BURNS_DIALOGUE.md"),),
    ),
    ChapterSpec(
        "03",
        "Chapter 3",
        tuple(SourceSpec(f"Beat {n}", f"CH03_B{n:02d}_*_DIALOGUE.md") for n in range(1, 16))
        + (
            SourceSpec("C06 — Nimera Takes Over a Table", "C06_NIMERA_TAKES_OVER_A_TABLE_DIALOGUE.md"),
            SourceSpec("C07 — Ilyra and Nimera", "C07_ILYRA_AND_NIMERA_DIALOGUE.md"),
        ),
    ),
)

DIALOGUE_RE = re.compile(r"^\*\*([^*\n]+):\*\*\s*(.*)$")
HEADING_RE = re.compile(r"^(#{2,6})\s+(.*)$")

SKIP_SECTION_TERMS = (
    "production note",
    "production notes",
    "person-brain / ensemble check",
    "person-brain / role-balance check",
    "audit check",
    "audit note",
    "knowledge check",
    "knowledge checkpoint",
    "world-state checkpoint",
    "continuity note",
    "continuity notes",
    "current production state",
    "conflict order",
    "gameplay handoff",
    "player exit",
    "cleanup window",
    "natural-turn",
    "spoken-dialogue",
    "presentation / canon check",
)

SKIP_LINE_PREFIXES = (
    "**Status:**",
    "**Story authority:**",
    "**Primary story authority:**",
    "**Method note:**",
    "**Rhythm note:**",
    "**Production note:**",
    "**Hard continuity:**",
    "**Owning cumulative manuscript:**",
    "**Dialogue workflow authority:**",
    "**Character-Life numbering authority:**",
    "**Availability:**",
    "**Participants:**",
    "**Objective:**",
    "Objective:",
    "Transition into:",
)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def plain_markdown(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^>\s*", "", text)
    return text.replace("**", "").replace("`", "").replace("*", "").strip()


def section_should_skip(heading: str) -> bool:
    low = plain_markdown(heading).lower()
    return any(term in low for term in SKIP_SECTION_TERMS)


def reader_blocks(source_text: str) -> list[tuple[str, str, str | None]]:
    """Extract player/reader-facing blocks while excluding writer-facing audit prose."""
    blocks: list[tuple[str, str, str | None]] = []
    skip_section = False
    in_fence = False

    for raw in source_text.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue

        heading = HEADING_RE.match(line)
        if heading:
            title = heading.group(2).strip()
            skip_section = section_should_skip(title)
            if not skip_section:
                cleaned = plain_markdown(title).strip("[]")
                if cleaned:
                    blocks.append(("heading", cleaned, None))
            continue

        if skip_section or not line.strip() or line.startswith("# "):
            continue
        if any(line.startswith(prefix) for prefix in SKIP_LINE_PREFIXES):
            continue
        if line.startswith("---") or line.startswith("|"):
            continue
        if line.startswith("-") and not line.startswith("—"):
            continue

        match = DIALOGUE_RE.match(line)
        if match:
            blocks.append(("dialogue", match.group(2).strip(), match.group(1).strip()))
            continue

        stripped = plain_markdown(line)
        if not stripped:
            continue
        upper = stripped.upper()
        if upper.startswith("OBJECTIVE:") or upper.startswith("CURRENT STATE:"):
            continue
        if "DO NOT " in upper and ("REVEAL" in upper or "ADD" in upper or "IDENTIFY" in upper):
            continue
        blocks.append(("prose", stripped, None))

    return blocks


def resolve_source(chapter: ChapterSpec, spec: SourceSpec) -> Path:
    matches = [p for p in chapter.directory.glob(spec.pattern) if p.is_file()]
    matches = [p for p in matches if p != chapter.output and "AUTHORITY_INDEX" not in p.name]
    if len(matches) != 1:
        rels = [str(p.relative_to(ROOT)) for p in matches]
        raise RuntimeError(
            f"{chapter.title} / {spec.label}: {spec.pattern!r} must resolve exactly once; "
            f"found {len(matches)}: {rels}"
        )
    return matches[0]


def resolved_chapter_sources(chapter: ChapterSpec) -> list[tuple[SourceSpec, Path, bytes, str]]:
    out: list[tuple[SourceSpec, Path, bytes, str]] = []
    seen: set[Path] = set()
    for spec in chapter.sources:
        path = resolve_source(chapter, spec)
        if path in seen:
            raise RuntimeError(f"Duplicate atomic source selected: {path}")
        seen.add(path)
        data = path.read_bytes()
        data.decode("utf-8")
        out.append((spec, path, data, sha256_bytes(data)))
    return out


def visible_chapter_text(sources: list[tuple[SourceSpec, Path, bytes, str]]) -> str:
    chunks: list[str] = []
    for _spec, _path, data, _digest in sources:
        for kind, text, speaker in reader_blocks(data.decode("utf-8")):
            chunks.append(f"{speaker}: {text}" if kind == "dialogue" else text)
    return "\n".join(chunks)


def verify_source_closure() -> None:
    if not SOURCE_CLOSURE.exists():
        raise RuntimeError(f"Missing source-closure record: {SOURCE_CLOSURE.relative_to(ROOT)}")
    closure = SOURCE_CLOSURE.read_text(encoding="utf-8")
    if "Chapters 0–3 are now fully updated at the source level" not in closure:
        raise RuntimeError("Source-closure record does not contain the required Chapters 0–3 closure decision.")

    for chapter in CHAPTERS:
        if not chapter.index.exists():
            raise RuntimeError(f"Missing chapter authority index: {chapter.index.relative_to(ROOT)}")
        index_text = chapter.index.read_text(encoding="utf-8")
        closure_markers = (
            "FULL SOURCE-LEVEL DIALOGUE CLOSURE",
            "LOCKED CURRENT CHAPTER-1 DIALOGUE",
            "CURRENT 12-BEAT ATOMIC DIALOGUE AUTHORITY",
        )
        if not any(marker in index_text for marker in closure_markers):
            raise RuntimeError(
                f"Chapter {int(chapter.chapter)} is not marked with a recognized source-closure status."
            )


def protected_checks(all_sources: dict[str, list[tuple[SourceSpec, Path, bytes, str]]]) -> None:
    raw_by_ch = {
        ch: "\n".join(data.decode("utf-8") for _, _, data, _ in sources)
        for ch, sources in all_sources.items()
    }

    must_exist = {
        "01": ("**CYANIS:** Old slut?", "**TORREN:** Bitch."),
        "03": (
            "PREVIOUS ERROR",
            "LAST SENTINEL CONFIRMED",
            "**CYANIS:** I bet you use that cape to sneak up on the goats you fuck.",
            "**TORREN:** You look like a walking dick in armor.",
            "Might. Elements. Grace. Memory. Perception. Ruin.",
        ),
    }
    for chapter, needles in must_exist.items():
        for needle in needles:
            if needle not in raw_by_ch[chapter]:
                raise RuntimeError(f"Protected anchor missing in Chapter {int(chapter)}: {needle}")

    stale_needles = (
        "Might. Elements. Grace. Resource. Perception. Ruin.",
        "Crest Arcanist",
        "Sixfold Knight",
    )
    for needle in stale_needles:
        for chapter, raw in raw_by_ch.items():
            if needle in raw:
                raise RuntimeError(f"Retired dialogue canon present in Chapter {int(chapter)} atomics: {needle}")

    ch1_visible = visible_chapter_text(all_sources["01"])
    if re.search(r"\bNimera\b", ch1_visible, flags=re.IGNORECASE):
        raise RuntimeError("Reader-visible Chapter 1 still references Nimera before her Chapter-3 meeting.")


def render_combined(chapter: ChapterSpec, sources: list[tuple[SourceSpec, Path, bytes, str]]) -> str:
    lines: list[str] = [
        f"# DIYSE — {chapter.title} — Synchronized Rehearsal-First Working Dialogue Manuscript",
        "",
        "**Status:** CURRENT DERIVED READ-THROUGH — SYNCHRONIZED AGAINST CURRENT ATOMIC DIALOGUE SOURCES",
        "",
        "**Authority rule:** standalone atomic scene files remain exact wording authority. This file is generated, not hand-authored. After any atomic edit, rerun `python tools/dialogue/sync_current_dialogue.py` before describing this manuscript as current.",
        "",
        "## Source manifest",
        "",
        "| Order | Canonical slot | Atomic source | SHA-256 |",
        "|---:|---|---|---|",
    ]
    for i, (spec, path, _data, digest) in enumerate(sources, start=1):
        lines.append(f"| {i} | {spec.label} | `{path.name}` | `{digest}` |")
    lines += ["", "---", ""]

    for spec, path, data, digest in sources:
        lines.extend([
            f"## {spec.label}",
            "",
            f"**Atomic source:** `{path.name}`  ",
            f"**Source SHA-256:** `{digest}`",
            "",
            data.decode("utf-8").rstrip(),
            "",
            "---",
            "",
        ])

    lines.extend([
        "## Synchronization footer",
        "",
        "This derived manuscript was assembled exclusively from the atomics listed above. No dialogue wording was rewritten during assembly.",
        "",
    ])
    return "\n".join(lines)


def render_sync_manifest(
    all_sources: dict[str, list[tuple[SourceSpec, Path, bytes, str]]],
    combined: dict[str, str],
) -> str:
    lines = [
        "# DIYSE — Chapters 0–3 Dialogue Synchronization Manifest",
        "",
        "**Status:** CURRENT — generated from source-closed active atomic dialogue",
        "",
        "The atomics remain exact wording authority. Combined manuscripts and the reader are generated derivatives.",
        "",
        "**Source closure:** `CHAPTER_00_03_FULL_SOURCE_CLOSURE_2026-09-13.md`",
        "",
        "## Combined manuscripts",
        "",
        "| Chapter | Combined manuscript | Combined SHA-256 | Atomic sources |",
        "|---|---|---|---:|",
    ]
    for chapter in CHAPTERS:
        lines.append(
            f"| {int(chapter.chapter)} | `{chapter.output.relative_to(ROOT)}` | "
            f"`{sha256_bytes(combined[chapter.chapter].encode('utf-8'))}` | "
            f"{len(all_sources[chapter.chapter])} |"
        )

    for chapter in CHAPTERS:
        lines.extend([
            "",
            f"## Chapter {int(chapter.chapter)} sources",
            "",
            "| Order | Canonical slot | Path | SHA-256 |",
            "|---:|---|---|---|",
        ])
        for i, (spec, path, _data, digest) in enumerate(all_sources[chapter.chapter], start=1):
            lines.append(f"| {i} | {spec.label} | `{path.relative_to(ROOT)}` | `{digest}` |")

    lines.extend([
        "",
        "## Verification rule",
        "",
        "Run `python tools/dialogue/sync_current_dialogue.py --check`. A non-zero exit means a combined manuscript or this manifest is stale relative to current atomics.",
        "",
    ])
    return "\n".join(lines)


def add_reader_docx(all_sources: dict[str, list[tuple[SourceSpec, Path, bytes, str]]]) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is required to build the reader") from exc

    READER_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(26)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(20)
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 3"].font.name = "Aptos"
    styles["Heading 3"].font.size = Pt(11.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DIYSE\nChapters 0–3")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Spoiler-Free Exact-Dialogue Reader")
    sr.bold = True
    sr.font.size = Pt(15)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run(
        "Current atomic dialogue edition. Spoken lines are copied verbatim from active production sources; "
        "writer-facing audits, implementation notes, and future-facing production metadata are omitted."
    )
    nr.italic = True
    doc.add_page_break()

    for chapter_i, chapter in enumerate(CHAPTERS):
        if chapter_i:
            doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_heading(f"Chapter {int(chapter.chapter)}", level=1)

        for spec, _path, data, _digest in all_sources[chapter.chapter]:
            doc.add_heading(spec.label, level=2)
            last_was_heading = False
            for kind, text, speaker in reader_blocks(data.decode("utf-8")):
                if kind == "heading":
                    if text.upper().startswith(("BOSS —", "BATTLE —")):
                        p = doc.add_paragraph()
                        r = p.add_run(text)
                        r.bold = True
                        r.italic = True
                    else:
                        doc.add_heading(text.title() if text.isupper() else text, level=3)
                    last_was_heading = True
                elif kind == "dialogue":
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(3)
                    who = p.add_run(f"{speaker}: ")
                    who.bold = True
                    p.add_run(text)
                    last_was_heading = False
                else:
                    p = doc.add_paragraph(text)
                    p.paragraph_format.space_after = Pt(4)
                    last_was_heading = False
            if not last_was_heading:
                doc.add_paragraph()

    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("DIYSE Chapters 0–3 — derived reader; atomic dialogue files remain authority")
    fr.font.size = Pt(8)

    doc.save(READER_OUT)


def prepare_sources() -> dict[str, list[tuple[SourceSpec, Path, bytes, str]]]:
    verify_source_closure()
    all_sources = {chapter.chapter: resolved_chapter_sources(chapter) for chapter in CHAPTERS}
    protected_checks(all_sources)
    return all_sources


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source-check",
        action="store_true",
        help="Verify source closure/source selection/protected anchors without checking derived outputs",
    )
    parser.add_argument("--check", action="store_true", help="Verify Markdown derived outputs without writing")
    parser.add_argument("--no-docx", action="store_true", help="Do not build the reader DOCX")
    args = parser.parse_args()

    if args.source_check and args.check:
        parser.error("--source-check and --check are mutually exclusive")

    all_sources = prepare_sources()

    if args.source_check:
        print("Chapters 0–3 source closure verified: source selection and protected guards pass.")
        for chapter in CHAPTERS:
            print(f"- Chapter {int(chapter.chapter)}: {len(all_sources[chapter.chapter])} atomic sources")
        return 0

    combined = {
        chapter.chapter: render_combined(chapter, all_sources[chapter.chapter]) for chapter in CHAPTERS
    }
    manifest = render_sync_manifest(all_sources, combined)
    expected: list[tuple[Path, str]] = [
        (chapter.output, combined[chapter.chapter]) for chapter in CHAPTERS
    ] + [(SYNC_MANIFEST, manifest)]

    if args.check:
        stale: list[str] = []
        for path, wanted in expected:
            got = path.read_text(encoding="utf-8") if path.exists() else ""
            if got != wanted:
                stale.append(str(path.relative_to(ROOT)))
        if stale:
            print("STALE derived dialogue outputs:")
            for path in stale:
                print(f"- {path}")
            return 1
        print("Dialogue combined manuscripts and sync manifest are current.")
        return 0

    for path, text in expected:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        print(f"wrote {path.relative_to(ROOT)}")

    if not args.no_docx:
        add_reader_docx(all_sources)
        print(f"wrote {READER_OUT.relative_to(ROOT)}")

    for path, wanted in expected:
        if path.read_text(encoding="utf-8") != wanted:
            raise RuntimeError(f"Post-write verification failed for {path}")

    print("Synchronization complete. Source authority indexes were not modified.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

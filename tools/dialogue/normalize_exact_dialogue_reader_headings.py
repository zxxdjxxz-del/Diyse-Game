#!/usr/bin/env python3
"""Normalize remaining source-coded reader headings without touching spoken dialogue."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
EXPECTED_DIALOGUE_LINES = 2250
LABEL_RE = re.compile(r"^.+:\s*$")
CODE_ONLY_RE = re.compile(r"^(?:P\d+|Beat\s+\d+)$", re.IGNORECASE)
CODE_PREFIX_RE = re.compile(r"^[CH]\d+\s*[—-]\s*(.+)$", re.IGNORECASE)


def is_dialogue(paragraph) -> bool:
    if not paragraph.runs:
        return False
    first = paragraph.runs[0]
    if not first.bold or not LABEL_RE.match(first.text):
        return False
    label = first.text.strip()[:-1].strip()
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def is_heading(paragraph) -> bool:
    return bool(paragraph.style and paragraph.style.name and paragraph.style.name.startswith("Heading"))


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def replace_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main() -> int:
    if not READER.exists():
        raise RuntimeError(f"Generated reader does not exist: {READER.relative_to(ROOT)}")

    document = Document(READER)
    before = [p.text for p in document.paragraphs if is_dialogue(p)]
    if len(before) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(f"Reader has {len(before)} spoken lines; expected {EXPECTED_DIALOGUE_LINES}.")

    changed = 0
    for paragraph in list(document.paragraphs):
        if not is_heading(paragraph):
            continue
        text = paragraph.text.strip()
        if CODE_ONLY_RE.fullmatch(text):
            remove_paragraph(paragraph)
            changed += 1
            continue
        match = CODE_PREFIX_RE.match(text)
        if match:
            replace_text(paragraph, match.group(1).strip())
            changed += 1

    after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if before != after or len(after) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError("Heading normalization altered spoken dialogue.")

    remaining = [
        p.text.strip()
        for p in document.paragraphs
        if is_heading(p) and (CODE_ONLY_RE.fullmatch(p.text.strip()) or CODE_PREFIX_RE.match(p.text.strip()))
    ]
    if remaining:
        raise RuntimeError(f"Coded source headings remain in reader: {remaining[:12]}")

    document.save(READER)
    print(f"Reader heading normalization complete; {changed} coded heading(s) normalized and {len(after)} spoken lines preserved exactly.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

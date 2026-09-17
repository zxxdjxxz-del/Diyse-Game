#!/usr/bin/env python3
"""Strip unmistakable production/meta residue from the generated exact-dialogue reader.

This is intentionally conservative: it removes only non-dialogue paragraphs/headings that
are plainly writer/gameplay-facing. Spoken dialogue is protected by an exact before/after
sequence assertion.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
DIALOGUE_PREFIX_RE = re.compile(r"^.+:\s*$")

META_HEADING_PATTERNS = (
    re.compile(r"^P\d{2}\s+end state$", re.I),
    re.compile(r"^Battle Character$", re.I),
    re.compile(r"^Encounter-direction priorities:?$", re.I),
    re.compile(r"^System / Party State$", re.I),
)

META_PROSE_PATTERNS = (
    re.compile(r"^Writer-facing identity note:", re.I),
    re.compile(r"^Hard scene rule:?$", re.I),
    re.compile(r"\bcurrent Chapter[- ]?\d+ story specifically authorizes\b", re.I),
    re.compile(r"\bencounter-authorized\b", re.I),
    re.compile(r"^This is the single combined final boss encounter\b", re.I),
    re.compile(r"^There is no argument for the sake of argument\b", re.I),
    re.compile(r"^Ilyra does not seize control of his care\b", re.I),
    re.compile(r"^Boss encounter begins\.?$", re.I),
    re.compile(r"^(?:Combat|Active battle|Permanent combat) party:\s*", re.I),
    re.compile(r"^TORREN HARTH\s+[—-]\s+PERMANENT PARTY MEMBER$", re.I),
    re.compile(r"^Cyanis \+ Ilyra \+ Torren$", re.I),
    re.compile(r"^Maevra remains a non-combat traveling companion\.?$", re.I),
    re.compile(r"^Beat\s+\d+\s+(?:ends|begins)\b", re.I),
    re.compile(r"^Hard story rule:?$", re.I),
    re.compile(r"^Encounter-direction priorities:?$", re.I),
    re.compile(r"^Use one restrained\b", re.I),
    re.compile(r"\bnot automatic control of the decision\b", re.I),
    re.compile(r"^The First Command Warden is the major .* confrontation of Chapter 3\.?$", re.I),
    re.compile(r"^Cyanis holds still because he wants to, not because Ilyra has taken control of him\.?$", re.I),
    re.compile(r"^Do not inspect the room yet\.?$", re.I),
    re.compile(r"^No playable road dungeon, encounter stretch, or travel-dialogue sequence occurs here\.?$", re.I),
    re.compile(r"^Cut directly to the party's arrival\b", re.I),
    re.compile(r"^Beat\s+15\s+[—-]\s+Cresthaven Headquarters / Cleanup Window\.?$", re.I),
    re.compile(r"^The road journey is skipped\.?$", re.I),
    re.compile(r"^The Card remains inert for the entire encounter\.?$", re.I),
    re.compile(r"^No added transformation or hidden meaning\.?$", re.I),
    re.compile(r"^This is not a defeat-to-death sequence and not a chase setup\.?$", re.I),
    re.compile(r"^The remaining Black Host presence does not mount another boss encounter\.?$", re.I),
    re.compile(r"^The party defeats Rhazek and wins the Old Bastion locally\..*not because the story reverses the victory\.?$", re.I),
    re.compile(r"^Optional short battle barks should stay practical and rare\.?$", re.I),
    re.compile(r"^Its battle identity should feel\b", re.I),
)


def is_heading(paragraph) -> bool:
    return bool(paragraph.style and paragraph.style.name and paragraph.style.name.startswith("Heading"))


def is_dialogue(paragraph) -> bool:
    if not paragraph.runs:
        return False
    first = paragraph.runs[0]
    if not first.bold or not DIALOGUE_PREFIX_RE.match(first.text):
        return False
    label = first.text.strip()[:-1].strip()
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def matches_any(text: str, patterns) -> bool:
    return any(pattern.search(text.strip()) for pattern in patterns)


def main() -> int:
    if not READER.exists():
        raise RuntimeError(f"Generated reader does not exist: {READER.relative_to(ROOT)}")

    document = Document(READER)
    dialogue_before = [p.text for p in document.paragraphs if is_dialogue(p)]
    removed = 0

    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if not text or is_dialogue(paragraph):
            continue
        if is_heading(paragraph) and matches_any(text, META_HEADING_PATTERNS):
            remove_paragraph(paragraph)
            removed += 1
            continue
        if matches_any(text, META_PROSE_PATTERNS):
            remove_paragraph(paragraph)
            removed += 1

    dialogue_after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if dialogue_before != dialogue_after:
        raise RuntimeError(
            "Meta-residue cleanup changed spoken dialogue; refusing to save "
            f"({len(dialogue_before)} before vs {len(dialogue_after)} after)."
        )

    document.save(READER)
    print(
        f"Reader meta-residue cleanup complete; removed {removed} non-dialogue paragraphs and preserved "
        f"{len(dialogue_after)} spoken dialogue lines exactly."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

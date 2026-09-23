#!/usr/bin/env python3
"""Remove the last source-code headings and writer-room narration from the Chapters 0-3 reader.

Runs after final presentation normalization and before layout repair. Spoken dialogue is
immutable: all 2,252 ALL-CAPS-speaker lines must remain byte-for-byte identical and in
exact order.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
EXPECTED_DIALOGUE_LINES = 2252
LABEL_RE = re.compile(r"^.+:\s*$")

PROSE_REPLACEMENTS = {
    "Dunmere is not discovered here. It was already the party's next practical destination.":
        "Dunmere remains the party's next destination.",
    "Ilyra Amarin is already working when Cyanis arrives. She was part of the wider convoy medical/Warden support, but she and Cyanis do not begin as established acquaintances.":
        "Ilyra Amarin is already working when Cyanis arrives, one of the convoy's medical/Warden support.",
    "There is no argument for the sake of argument. He got the information he needed; she sees that he will actually use it.":
        "He got the information he needed; she sees that he will actually use it.",
    "This is the first incomplete Card response.": "The Card gives an incomplete response.",
    "Nobody infers who he is, why he withdrew, or whether the north cut was bait.":
        "They are left without answers about who he was or why he withdrew.",
    "The fight does not pause.": "The fighting continues.",
    "The immediate organized Black Host attack has failed, but the convoy crisis is not treated as a victory celebration.":
        "The organized Black Host attack has failed, but the convoy is still in crisis.",
    "The Broken Convoy does not celebrate. The immediate organized Black Host push has failed, but the field still contains wounded people, missing people, damaged wagons, and unfinished work.":
        "The field still contains wounded people, missing people, damaged wagons, and unfinished work.",
    "No one turns this into a discussion of what the Card is. Torren receives only the travel-relevant facts.":
        "Torren is told only what matters for the road.",
    "Nobody promotes that timing into motive.": "The timing is noted without assigning a motive.",
    "It is a small ordinary exchange, not an attempt to manage Torren's grief.":
        "The exchange stays small and ordinary.",
    "Torren points out only the practical relationship, not an invented purpose.":
        "Torren points out the practical relationship between the structures.",
    "That is enough. Nobody names the ancient system or settles what the passages were for.":
        "They still cannot name the ancient system or settle what the passages were for.",
    "The Card does not speak, open, manifest a Prime, reveal a Face, identify Cyanis, or provide explanatory text.":
        "The Card remains inert.",
    "Nobody explains why the Castellan activated.": "Why the Castellan activated remains unknown.",
    "Nobody gives the monument a name or function.": "The monument's name and function remain unknown.",
    "Cyanis does not make Torren repeat the decision in ceremonial language.": "Cyanis accepts Torren's decision.",
    "Nobody names the monument or assigns it a function.": "The monument remains unnamed, its function unknown.",
    "Nobody names the Cistern Devourer or knows what is waiting beyond that access.":
        "What waits beyond that access is unknown.",
    "The damaged edge remains unresolved. Nobody invents the missing information.":
        "The damaged edge remains unresolved.",
    "That is enough to turn Torren from social-short into route-long.":
        "Torren grows more expansive once the route becomes the subject.",
    "Torren answers the whole tactical point instead of forcing Cyanis to pull it out one clause at a time.":
        "Torren gives the full tactical answer.",
    "Nobody explains what the crest-marked structure is or why the Crest appears there.":
        "The crest-marked structure's purpose remains unknown.",
    "Torren owns route relationships. Ilyra captures stable visual markers, the huge city shape, and the crest-marked northern structure. Cyanis keeps the Wayfinder overlap straight. Maevra checks that the copy preserves what they actually saw rather than what they think it means.":
        "Torren works through the route relationships. Ilyra copies stable visual markers, the huge city shape, and the crest-marked northern structure. Cyanis keeps the Wayfinder overlap straight. Maevra checks the copy against what they actually saw.",
    "Nobody asks the prisoners to manufacture an onward destination they do not know.":
        "The prisoners do not know an onward destination.",
    "That is the problem in one line; nobody restates it three more ways.": "The problem is clear.",
    "The sentry immediately raises the alarm rather than delivering a villain speech.":
        "The sentry immediately raises the alarm.",
    "Movement changes around them. Doors shut. Boots converge on useful defensive positions rather than appearing from nowhere.":
        "Movement changes around them. Doors shut. Boots converge on defensive positions.",
    "The Host response is organized but finite. Existing soldiers reposition toward chokepoints, upper stairs, and command access. The fort does not generate endless reinforcements. Support personnel move out of the fighting where plausible; wounded soldiers are not used as shields or disposable obstacles.":
        "The Host response is organized. Soldiers reposition toward chokepoints, upper stairs, and command access while support personnel and wounded move clear of the fighting.",
    "He does not reach for theatrics or pretend surprise.": "He shows no surprise.",
    "The masked officer's attention shifts to the trio. He remains ready, but does not move to begin a separate confrontation.":
        "The masked officer's attention shifts to the trio. He remains ready but holds position.",
    "He remains ready but does not move to begin a separate confrontation.": "He remains ready but holds position.",
    "Dunmere is functioning around the return rather than staging a formal reception.": "Dunmere keeps moving around the return.",
    "Rhazek fights as the commander of a functioning position rather than as a theatrical duelist waiting for a heroic showdown.":
        "Rhazek fights from the command position, directing a functioning defense.",
    "Ilyra stays alert rather than automatically converting the standoff into a post-battle medical check.":
        "Ilyra stays alert through the standoff.",
    "Torren checks the maintained road-side exit and its route rather than beginning another dungeon mechanism sequence.":
        "Torren checks the maintained road-side exit and its route.",
    "Dunmere is functioning around the return rather than staging a formal reception. Rescued people are being reunited, seated, fed, checked, or helped farther into town as needed.":
        "Dunmere keeps moving around the return. Rescued people are reunited, seated, fed, checked, or helped farther into town as needed.",
    "For a little while they complain about the cold instead. That conversation is not important enough to transcribe every line.":
        "For a little while, they complain about the cold instead.",
    "Lysara accepts the boundary without turning it into a diagnosis.": "Lysara accepts the boundary.",
    "No one accuses the seal staff. No one proposes testing the Card against the seal.":
        "The seal staff are not accused, and the Card remains untouched.",
    "The Scribe Engine is disabled rather than destroyed spectacularly. Its active arms lose tension and the local hostile motion stops. The deeper Archive remains active around it.":
        "The Scribe Engine is disabled. Its active arms lose tension and the local hostile motion stops. The deeper Archive remains active around it.",
    "Nimera examines only what remains legible without pretending she can reconstruct the machine's entire historical purpose.":
        "Nimera examines what remains legible.",
    "Cyanis accepts the decision without turning it into a ceremony.": "Cyanis accepts the decision.",
    "No one attempts to interpret Last Sentinel yet.": "The meaning of Last Sentinel remains unresolved.",
    "Nobody names what Last Sentinel refers to or claims the timing proves a mechanism.":
        "What Last Sentinel refers to, and why the timing matters, remains unresolved.",
    "Mirena does not narrate each function.": "Mirena moves through the work without commentary.",
    "No one names the Archive Judgment Engine.": "The machine remains unnamed.",
    "Maevra reads the operational implication rather than summarizing every record.":
        "Maevra reads the operational implication.",
}

PROSE_DROPS = {
    "The Card does not speak, identify Cyanis, reveal a Face, manifest a Prime, or explain the activation.",
    "The earlier side access is now available for the optional Cistern Devourer Hunt.",
    "Dunmere is not discovered here. It was already the party's intended destination.",
    "No one stops to restate the entire Junction discovery.",
    "No one revisits LAST SENTINEL CONFIRMED, the seal copies, or Cresthaven's Ancient mystery.",
}

FORBIDDEN_PHRASES = (
    "established acquaintances",
    "argument for the sake of argument",
    "first incomplete Card response",
    "Nobody infers",
    "fight does not pause",
    "not treated as a victory celebration",
    "travel-relevant facts",
    "promotes that timing into motive",
    "not an attempt to manage",
    "invented purpose",
    "That is enough. Nobody names",
    "manifest a Prime",
    "reveal a Face",
    "Nobody explains why",
    "Nobody gives the monument",
    "ceremonial language",
    "Cistern Devourer",
    "Nobody invents",
    "not discovered here",
    "restate the entire Junction",
    "social-short",
    "whole tactical point instead of forcing",
    "Torren owns route relationships",
    "manufacture an onward destination",
    "nobody restates it",
    "villain speech",
    "appearing from nowhere",
    "endless reinforcements",
    "reach for theatrics",
    "separate confrontation",
    "theatrical duelist",
    "post-battle medical check",
    "dungeon mechanism",
    "staging a formal reception",
    "not important enough to transcribe",
    "turning it into a diagnosis",
    "proposes testing the Card",
    "destroyed spectacularly",
    "pretending she can reconstruct",
    "turning it into a ceremony",
    "attempts to interpret Last Sentinel",
    "claims the timing proves",
    "revisits LAST SENTINEL",
    "does not narrate each function",
    "Archive Judgment Engine",
)


def is_dialogue(paragraph) -> bool:
    if not paragraph.runs:
        return False
    first = paragraph.runs[0]
    if not first.bold or not LABEL_RE.match(first.text):
        return False
    label = first.text.strip()[:-1].strip()
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def replace_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def main() -> int:
    if not READER.exists():
        raise RuntimeError(f"Generated reader does not exist: {READER.relative_to(ROOT)}")

    document = Document(READER)
    dialogue_before = [p.text for p in document.paragraphs if is_dialogue(p)]
    if len(dialogue_before) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(
            f"Reader has {len(dialogue_before)} spoken lines; expected {EXPECTED_DIALOGUE_LINES}."
        )

    heading_changes = 0
    prose_changes = 0
    for paragraph in list(document.paragraphs):
        if is_dialogue(paragraph):
            continue

        if paragraph.style and paragraph.style.name == "Heading 2":
            text = paragraph.text.strip()
            if re.fullmatch(r"(?:P\d+|Beat\s+\d+)", text, re.IGNORECASE):
                remove_paragraph(paragraph)
                heading_changes += 1
                continue
            match = re.match(r"^[CH]\d+\s*[—-]\s*(.+)$", text, re.IGNORECASE)
            if match:
                replace_text(paragraph, match.group(1).strip())
                heading_changes += 1
                continue

        text = paragraph.text.strip()
        if text in PROSE_DROPS:
            remove_paragraph(paragraph)
            prose_changes += 1
            continue
        replacement = PROSE_REPLACEMENTS.get(text)
        if replacement is not None:
            replace_text(paragraph, replacement)
            prose_changes += 1

    dialogue_after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if dialogue_before != dialogue_after or len(dialogue_after) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError("Story-facing reader polish altered spoken dialogue.")

    coded_headings = [
        p.text.strip()
        for p in document.paragraphs
        if p.style
        and p.style.name == "Heading 2"
        and (
            re.fullmatch(r"(?:P\d+|Beat\s+\d+)", p.text.strip(), re.IGNORECASE)
            or re.match(r"^[CH]\d+\s*[—-]", p.text.strip(), re.IGNORECASE)
        )
    ]
    if coded_headings:
        raise RuntimeError(f"Coded source headings remain in reader: {coded_headings[:12]}")

    residue = []
    for paragraph in document.paragraphs:
        if is_dialogue(paragraph):
            continue
        text = paragraph.text.strip()
        for phrase in FORBIDDEN_PHRASES:
            if phrase.casefold() in text.casefold():
                residue.append(text)
                break
    if residue:
        preview = "\n".join(f"- {line}" for line in residue[:12])
        raise RuntimeError(f"Writer-room narration remains after story polish:\n{preview}")

    document.save(READER)
    print(
        f"Story-facing reader polish complete; {heading_changes} source heading(s) normalized, "
        f"{prose_changes} narration paragraph(s) normalized, and {len(dialogue_after)} spoken lines preserved exactly."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

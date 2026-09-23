#!/usr/bin/env python3
"""Final story-facing normalization for the generated Chapters 0-3 reader.

Runs after clean_exact_dialogue_reader.py. It removes or rewrites the last known
production-language fragments and verifies that all spoken dialogue remains byte-for-byte
identical in order and wording.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
EXPECTED_DIALOGUE_LINES = 2103
LABEL_RE = re.compile(r"^.+:\s*$")

HEADING_REPLACEMENTS = {
    "working royal audience": "Royal Audience",
    "post-boss": "After the Battle",
}
FORBIDDEN_HEADING_PATTERNS = (
    re.compile(r"^working\s+royal\s+audience$", re.IGNORECASE),
    re.compile(r"^post[- ]boss$", re.IGNORECASE),
    re.compile(r"^c\d+\s+function$", re.IGNORECASE),
)

PROSE_REPLACEMENTS = {
    "Current atomic dialogue edition. Spoken lines are copied verbatim from the active source scenes; production notes are omitted.":
        "Current atomic dialogue edition. Spoken lines are copied verbatim from the active source material; non-story notes are omitted.",
    "From this moment onward, the recovery casing is broken beyond use and leaves active continuity.":
        "The recovery casing is broken beyond use.",
    "The casing is gone from continuity.": "The casing is gone.",
    "The extra second is not a professional beat. Ilyra has just seen a room full of dead people and reacts as a person before moving on.":
        "Ilyra has just seen a room full of dead people. She takes an extra second before moving on.",
    "All four current characters are present. Torren's dinner is genuinely bad but hot, filling, safe, and cheap. Maevra already knows his cooking well enough to have old complaints ready; Cyanis and Ilyra are discovering it.":
        "Torren's dinner is genuinely bad but hot, filling, safe, and cheap. Maevra already knows his cooking well enough to have old complaints ready; Cyanis and Ilyra are discovering it.",
    "The medical task gives them a natural reason to sit together. Ilyra's curiosity about Torren emerges from what she has actually watched all chapter: Maevra and Torren know one another extremely well, argue constantly, and anticipate each other's habits with suspicious ease.":
        "With the medical work underway, Ilyra's curiosity turns to Torren. Maevra and Torren know one another extremely well, argue constantly, and anticipate each other's habits with suspicious ease.",
    "The party steps slightly aside from wagon traffic. This is a short practical stop, not a strategy council.":
        "The party steps slightly aside from wagon traffic for a moment.",
    "After a short post-boss exploration stretch, the Ancient route narrows near a section where later construction presses against the older structure.":
        "After a short exploration stretch, the Ancient route narrows near a section where later construction presses against the older structure.",
    "This is a choice of priority, not inability. Rhazek has been beaten; the party chooses the rescue over turning a local victory into a chase.":
        "Rhazek has been beaten; the party turns back toward the prisoners instead of chasing him.",
    "The party comes back for them. The rescue is completed as a human action, not skipped because the boss is over.":
        "The party comes back for them.",
    "The exchange is brief. Work continues around them.": "Work continues around them.",
    "This is a working royal audience, not a ceremonial spectacle. Queen Lysara, Crown Princess Mirena, Chancellor Othmar Calder, Maevra, Cyanis, Ilyra, and Torren are present with only the clerical support needed for records and maps.":
        "Queen Lysara, Crown Princess Mirena, Chancellor Othmar Calder, Maevra, Cyanis, Ilyra, and Torren are present with only the clerical support needed for records and maps.",
    "This is a practical Crown workspace, not a sinister vault. Necessary staff produce records and the Queen's physical seal, then recede from the conversation.":
        "Necessary staff produce records and the Queen's physical seal, then recede from the room.",
    "This is not new evidence. It is the same old material already reported in Caelora:":
        "It is the same material already reported in Caelora:",
    "Movement pauses.": "The group stops.",
    "Movement pauses briefly.": "The group stops briefly.",
    "Movement pauses long enough for the space to register.": "The group stops long enough to take in the space.",
    "The party is given a clean visual read before anyone speaks.": "The party takes in the space before anyone speaks.",
    "Once ordinary wilderness traversal resumes, Cyanis is again the sole visible party field character.":
        "Once they return to the wilderness, Cyanis takes the lead again.",
    "At the established point where later construction visibly interfaces with older Diysean structure, movement pauses.":
        "Later construction visibly meets the older Diysean structure.",
    "At the channel carrying the Black Host breach farther inward, movement pauses and the relevant characters may appear again.":
        "At the channel carrying the Black Host breach farther inward, the group stops.",
    "At the mural viewing point, movement pauses.":
        "The group stops at the mural.",
    "At the established overlook, movement pauses.":
        "The group stops at the overlook.",
    "Torren traces the relevant relationship on his own map without turning it into a lecture for the audience.":
        "Torren traces the relationship on his own map.",
    "The elder indicates an old municipal route on a local map or equivalent field reference.":
        "The elder indicates an old municipal route on a local map.",
    "The elder gives the practical location of the old waterworks access. Exact map geometry and access mechanism belong to area design.":
        "The elder gives the location of the old waterworks access.",
    "The waterworks route opens into a much larger Ancient interior. Movement pauses long enough for a clean visual read.":
        "The waterworks route opens into a much larger Ancient interior. The group stops to take it in.",
    "Exploration reaches a chamber dominated by a large wall-scale map. The relevant group models appear at a natural stopping point.":
        "The group reaches a chamber dominated by a large wall-scale map.",
    "The area itself shows why the center route is unsafe. No second explanation is needed.":
        "The area itself shows why the center route is unsafe.",
    "The line is observational, not therapeutic. Ilyra does not ask him to unpack it.":
        "Ilyra does not ask him to unpack it.",
    "That is the end of it. No relationship explanation follows.":
        "They leave it there.",
    "The tension remains; nobody turns the moment into a lecture about the ruin.":
        "The tension remains, and nobody says more.",
    "The relief supplies the visible inventory. The cast speaks only to what it implies.":
        "Relief supplies fill the space.",
    "The mural is the one major visual asset for this beat.":
        "The mural dominates the chamber.",
    "The monument is clearly important, but no crowd descends beneath it and no deep opening is shown there.":
        "The monument is clearly important, but no crowd descends beneath it and there is no deep opening there.",
    "The group takes in the mural as a whole. The mural itself already establishes the many openings and civilian-scale movement, so nobody repeats the inventory from the lower junction.":
        "The group takes in the mural as a whole. Many openings and signs of civilian-scale movement are visible across it.",
    "The Briarhide is still not shown.":
        "The Briarhide remains out of sight.",
    "Torren does not smoke while eating. After everyone has finished, he lights a blunt as ordinary post-meal behavior.":
        "Torren waits until everyone has finished eating before lighting a blunt.",
    "Once ordinary wilderness traversal begins, Cyanis returns to being the sole visible party field character.":
        "Once they return to the wilderness, Cyanis takes the lead again.",
    "Then, preserving the exact locked exchange:":
        "Then:",
    "Ilyra notices the injury but does not turn first contact into forced treatment.":
        "Ilyra notices the injury but does not press treatment on him.",
    "The elder's expression hardens, but does not turn into a speech.":
        "The elder's expression hardens, but he says no more.",
    "The humor is brief and ordinary. It does not erase the unresolved missing people.":
        "The laugh passes quickly; the missing people still weigh on them.",
    "The joke is brief and the boundary remains Crown-owned, not Ilyra-owned.":
        "The joke passes; the Crown's boundary remains unchanged.",
    "The conversation ends before the room turns the departure into another briefing.":
        "The conversation ends there.",
    "She keeps the relevant citations rather than turning the discovery into a lecture.":
        "She keeps only the relevant citations.",
}

PROSE_DROPS = {
    "The exchange is about to end.",
    "The Card response is not selectable and is not a Prime activation.",
    "This is a genuine medical/travel question, so Ilyra owns the answer.",
    "This is boss animation owned by the encounter itself.",
    "Maevra remains a non-combat traveling companion.",
    "One pot and ordinary bowls are enough. No cooking montage or food-detail animation is needed.",
    "The opening stays quick because the bad-food reveal benefits from it. Then the meal settles.",
    "Encounter-direction priorities:",
    "That quick three-line escalation stays quick because the joke benefits from it.",
}

FORBIDDEN_PROSE_PATTERNS = (
    re.compile(r"\bnot selectable\b", re.IGNORECASE),
    re.compile(r"\bboss animation\b", re.IGNORECASE),
    re.compile(r"\bowns the answer\b", re.IGNORECASE),
    re.compile(r"\bpost-boss\b", re.IGNORECASE),
    re.compile(r"\bencounter-direction priorities\b", re.IGNORECASE),
    re.compile(r"\bcooking montage\b", re.IGNORECASE),
    re.compile(r"\bbenefits from it\b", re.IGNORECASE),
    re.compile(r"\ball chapter\b", re.IGNORECASE),
    re.compile(r"\bnot skipped because the boss\b", re.IGNORECASE),
    re.compile(r"\bactive continuity\b", re.IGNORECASE),
    re.compile(r"\bgone from continuity\b", re.IGNORECASE),
    re.compile(r"\bmovement pauses\b", re.IGNORECASE),
    re.compile(r"\bfor the audience\b", re.IGNORECASE),
    re.compile(r"\bequivalent field reference\b", re.IGNORECASE),
    re.compile(r"\bexact map geometry\b", re.IGNORECASE),
    re.compile(r"\barea design\b", re.IGNORECASE),
    re.compile(r"\bgroup models?\b", re.IGNORECASE),
    re.compile(r"\bnatural stopping point\b", re.IGNORECASE),
    re.compile(r"\bclean visual read\b", re.IGNORECASE),
    re.compile(r"\bno second explanation is needed\b", re.IGNORECASE),
    re.compile(r"\bthe line is observational\b", re.IGNORECASE),
    re.compile(r"\brelationship explanation follows\b", re.IGNORECASE),
    re.compile(r"\bthe cast speaks\b", re.IGNORECASE),
    re.compile(r"\bvisual asset\b", re.IGNORECASE),
    re.compile(r"\bis shown there\b", re.IGNORECASE),
    re.compile(r"\bvisible party field character\b", re.IGNORECASE),
    re.compile(r"\blocked exchange\b", re.IGNORECASE),
    re.compile(r"\bordinary post-meal behavior\b", re.IGNORECASE),
    re.compile(r"\bturn(?:s|ing) .* into a lecture\b", re.IGNORECASE),
    re.compile(r"\banother briefing\b", re.IGNORECASE),
)


def is_heading(paragraph) -> bool:
    return bool(paragraph.style and paragraph.style.name and paragraph.style.name.startswith("Heading"))


def is_dialogue(paragraph) -> bool:
    if not paragraph.runs:
        return False
    first = paragraph.runs[0]
    if not first.bold or not LABEL_RE.match(first.text):
        return False
    label = first.text.strip()[:-1].strip()
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def replace_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def main() -> int:
    if not READER.exists():
        raise RuntimeError(f"Generated reader does not exist: {READER.relative_to(ROOT)}")

    document = Document(READER)
    dialogue_before = [p.text for p in document.paragraphs if is_dialogue(p)]
    if len(dialogue_before) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(
            f"Reader has {len(dialogue_before)} spoken lines; expected {EXPECTED_DIALOGUE_LINES}."
        )

    heading_changes = 0
    prose_changes = 0
    for paragraph in list(document.paragraphs):
        if is_heading(paragraph):
            key = paragraph.text.strip().casefold()
            replacement = HEADING_REPLACEMENTS.get(key)
            if replacement is not None:
                replace_text(paragraph, replacement)
                heading_changes += 1
            continue

        if is_dialogue(paragraph):
            continue

        text = paragraph.text.strip()
        if text in PROSE_DROPS:
            remove_paragraph(paragraph)
            prose_changes += 1
            continue
        replacement = PROSE_REPLACEMENTS.get(text)
        if replacement is not None:
            replace_text(paragraph, replacement)
            prose_changes += 1

    dialogue_after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if dialogue_before != dialogue_after or len(dialogue_after) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError("Final reader normalization altered spoken dialogue.")

    forbidden_headings = [
        p.text.strip()
        for p in document.paragraphs
        if is_heading(p)
        and any(pattern.search(p.text.strip()) for pattern in FORBIDDEN_HEADING_PATTERNS)
    ]
    if forbidden_headings:
        raise RuntimeError(f"Forbidden production-facing reader headings remain: {forbidden_headings}")

    forbidden_prose = [
        p.text.strip()
        for p in document.paragraphs
        if not is_heading(p)
        and not is_dialogue(p)
        and any(pattern.search(p.text.strip()) for pattern in FORBIDDEN_PROSE_PATTERNS)
    ]
    if forbidden_prose:
        preview = "\n".join(f"- {line}" for line in forbidden_prose[:12])
        raise RuntimeError(f"Final reader production prose remains:\n{preview}")

    document.save(READER)
    print(
        f"Final reader normalization complete; {heading_changes} heading(s) and "
        f"{prose_changes} prose paragraph(s) normalized; {len(dialogue_after)} spoken lines preserved exactly."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Restore the approved world intro and currently validated encounter bridges to the Ch0-3 reader.

Run after presentation cleanup/heading normalization and before layout repair.
All 2,038 spoken lines are immutable.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
INTRO = ROOT / "docs/04_WORLD_AND_LORE/PLAYER_FACING_WORLD_INTRO.md"
CH0 = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/CHAPTER_ENEMIES/CHAPTER_00.md"
CH1 = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/CHAPTER_ENEMIES/CHAPTER_01.md"
CH2 = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/CHAPTER_ENEMIES/CHAPTER_02.md"
CH3 = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/CHAPTER_ENEMIES/CHAPTER_03.md"
CH1_FORMATIONS = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/ENCOUNTER_FORMATIONS/CHAPTER_01_FORMATIONS.md"
CH2_FORMATIONS = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/ENCOUNTER_FORMATIONS/CHAPTER_02_FORMATIONS.md"
CH3_FORMATIONS = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/ENCOUNTER_FORMATIONS/CHAPTER_03_FORMATIONS.md"
PLACEMENT = ROOT / "docs/09_ENEMIES_AND_ENCOUNTERS/CURRENT_PLACEMENT_AND_RESOLUTION_CORRECTIONS_2026-09-12.md"

EXPECTED_DIALOGUE_LINES = 2038
LABEL_RE = re.compile(r"^.+:\s*$")

# Chapter-level roster summaries are inserted immediately after each chapter heading.
# The body is reader-facing; the owning source and needles are validation-only.
ROSTER_SUMMARIES = (
    ("Chapter 0", "Current Enemy Roster", (
        "Ordinary / tutorial: Black Host Raider • Black Host Crossbowman • Ruin Shieldbearer • Convoy Rift Hound",
        "Authored protected encounter: Ruin Vanguard Pursuer",
        "Final boss: Riftmaw + Convoy War-Sorcerer",
    ), CH0, ("Black Host Raider", "Convoy Rift Hound", "Ruin Vanguard Pursuer", "Riftmaw", "Convoy War-Sorcerer")),
    ("Chapter 1", "Current Enemy Roster", (
        "Northern Briar: Greenhollow Stalker • Thornvine Creeper • Briar Boar",
        "Southern Briar adds: Needlewing • Rootmaw • Brambleback",
        "Hollow Watch — Black Host: Black Host Raider • Black Host Crossbowman • Ruin Shieldbearer",
        "Hollow Watch — constructs: Watch Sentry • Watch Ballista • Watch Captain Frame",
        "Mini-boss: Hollow Watch Castellan",
        "Main / final boss: Briarhide Stalker",
        "Regional Hunt #1: Cistern Devourer",
    ), CH1, ("Greenhollow Stalker", "Needlewing", "Watch Sentry", "Watch Captain Frame", "Hollow Watch Castellan", "Briarhide Stalker", "Cistern Devourer")),
    ("Chapter 2", "Current Enemy Roster", (
        "Old Waterworks: Bogshell • Cistern Leech • Needlewing",
        "Sunken Archive: Archive Current • Memory Scribe, with compatible Waterworks wildlife in flooded sectors",
        "Old Bastion: Black Host Raider • Black Host Crossbowman • Ruin Shieldbearer • Black Host War-Sorcerer • Rift Hound",
        "Bosses: Archive Leviathan • Commander Rhazek — Bastion Master",
        "Regional Hunt #2: Scaldback",
    ), CH2, ("Bogshell", "Archive Current", "Memory Scribe", "Black Host War-Sorcerer", "Rift Hound", "Archive Leviathan", "Commander Rhazek", "Scaldback")),
    ("Chapter 3", "Current Enemy Roster", (
        "Old City Archives: Judgment Frame • Erasure Wisp • Authority Lens • Archive Current",
        "Late Old City strong normal-pool Elite: Grand Inquisitor Frame",
        "Cresthaven Ancient tower base: Watch Sentry • Watch Ballista • Watch Captain Frame • Command Guard Frame • Authority Lens • Command Ring Drone",
        "Mandatory bosses: Archive Scribe Engine • First Command Warden",
        "Regional Hunt #3: Archive Judgment Engine",
        "No hostile Caelora / road pool; Way-Fort Marauder, Rift Boltman, and Black Host Ward-Sorcerer are retired from Chapter 3.",
    ), CH3, ("Judgment Frame", "Erasure Wisp", "Authority Lens", "Archive Current", "Grand Inquisitor Frame", "Command Guard Frame", "Command Ring Drone", "Archive Scribe Engine", "First Command Warden", "Archive Judgment Engine")),
)

# target heading, inserted heading, body paragraphs, owning source, required authority terms
BRIDGES = (
    ("Cyanis Solo", "Opening Ambush", (
        "Combat 1 — Cyanis solo: Black Host Raider + Black Host Crossbowman.",
        "Combat 2 — Cyanis solo: Black Host Raider + Ruin Shieldbearer.",
        "Combat 3 — Cyanis solo: 2 Convoy Rift Hounds.",
        "Chapter 0 uses authored/tutorial encounters rather than the normal random-encounter cadence.",
    ), CH0, ("Black Host Raider", "Black Host Crossbowman", "Ruin Shieldbearer", "2 Convoy Rift Hounds")),
    ("Hound Pressure", "Wreck Field", (
        "Combat 4 — Cyanis solo: Black Host Crossbowman + Convoy Rift Hound.",
        "Combat 5 — Cyanis solo: one lone Convoy Rift Hound threatening the survivor route.",
    ), CH0, ("Black Host Crossbowman + Convoy Rift Hound", "one lone Convoy Rift Hound")),
    ("The Pursuer", "Concealed Ruin Vanguard", (
        "Combat party: Cyanis + Ilyra.",
        "Enemy: Ruin Vanguard Pursuer. The pursuer's identity remains unknown to the party.",
    ), CH0, ("Ruin Vanguard Pursuer",)),
    ("Final Push", "Broken Convoy", (
        "Combat party: Cyanis + Ilyra.",
        "Boss encounter: Riftmaw + Convoy War-Sorcerer.",
    ), CH0, ("Riftmaw", "Convoy War-Sorcerer")),

    ("Halfway Stop", "Random Encounters — Northern Briar Passage", (
        "Greenhollow Stalker • Thornvine Creeper • Briar Boar",
    ), CH1, ("Greenhollow Stalker", "Thornvine Creeper", "Briar Boar")),
    ("Hollow Watch Reveal", "Random Encounters — Greenhollow / Hollow Watch Approach", (
        "Greenhollow Stalker • Thornvine Creeper • Briar Boar",
    ), CH1, ("Greenhollow Stalker", "Thornvine Creeper", "Briar Boar")),
    ("Garrison Discovery", "Random Encounters — Hollow Watch", (
        "Black Host Raider • Black Host Crossbowman • Ruin Shieldbearer • Watch Sentry • Watch Ballista • Watch Captain Frame",
    ), CH1, ("Black Host Raider", "Black Host Crossbowman", "Ruin Shieldbearer", "Watch Sentry", "Watch Ballista", "Watch Captain Frame")),
    ("First Clear Sighting", "Random Encounters — Southern Briar", (
        "Greenhollow Stalker • Thornvine Creeper • Briar Boar • Needlewing • Rootmaw • Brambleback",
    ), CH1, ("Greenhollow Stalker", "Thornvine Creeper", "Briar Boar", "Needlewing", "Rootmaw", "Brambleback")),

    ("Sealed Side Door", "Random Encounters — Old Waterworks", (
        "Bogshell • Cistern Leech • Needlewing",
    ), CH2, ("Bogshell", "Cistern Leech", "Needlewing")),
    ("Threshold", "Random Encounters — Sunken Archive", (
        "Archive Current • Memory Scribe • Bogshell • Cistern Leech • Needlewing",
    ), CH2, ("Archive Current", "Memory Scribe", "Bogshell", "Cistern Leech", "Needlewing")),
    ("The Alarm", "Random Encounters — Old Bastion", (
        "Ruin Shieldbearer • Black Host Crossbowman • Black Host War-Sorcerer • Black Host Raider • Rift Hound",
    ), CH2, ("Ruin Shieldbearer", "Black Host Crossbowman", "Black Host War-Sorcerer", "Black Host Raider", "Rift Hound")),

    ("Lower Archives", "Random Encounters — Old City Archives", (
        "Lower Archives through Hall of Seals: Judgment Frame • Erasure Wisp • Authority Lens.",
        "Deep Archives adds Archive Current. Grand Inquisitor Frame is a rare late strong normal-pool Elite, maximum one per formation.",
        "Archive Scribe Engine is the mandatory Beat-11 boss and never appears in the random pool.",
    ), CH3, ("Judgment Frame", "Erasure Wisp", "Authority Lens", "Archive Current", "Grand Inquisitor Frame", "Archive Scribe Engine")),
    ("Cresthaven / Ancient Tower Base / First Command Warden", "Random Encounters — Cresthaven Ancient Tower Base", (
        "Watch Sentry • Watch Ballista • Watch Captain Frame • Command Guard Frame • Authority Lens • Command Ring Drone",
        "Watch Captain Frame remains maximum one per formation. The immediate First Command Warden approach is safe.",
    ), CH3, ("Watch Sentry", "Watch Ballista", "Watch Captain Frame", "Command Guard Frame", "Authority Lens", "Command Ring Drone", "First Command Warden")),
)

FORMATION_AUTHORITIES = (
    (CH1_FORMATIONS, ("Briar Passage — first / northern traversal", "Hollow Watch — occupied fort / Black Host", "Southern Briar Passage", "Hollow Watch Castellan", "Briarhide Stalker")),
    (CH2_FORMATIONS, ("Old Waterworks", "Sunken Archive", "Old Bastion", "Watch Sentry in Chapter 2", "Full Bastion Response")),
    (CH3_FORMATIONS, ("Old City Archives", "Cresthaven Ancient tower base", "Archive Scribe Engine", "First Command Warden", "Way-Fort Marauder / Rift Boltman / Black Host Ward-Sorcerer are retired from Chapter 3")),
)


def is_dialogue(paragraph) -> bool:
    if not paragraph.runs:
        return False
    first = paragraph.runs[0]
    if not first.bold or not LABEL_RE.match(first.text):
        return False
    label = first.text.strip()[:-1].strip()
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def find_heading(document: Document, text: str):
    matches = [
        p for p in document.paragraphs
        if p.style and p.style.name.startswith("Heading") and p.text.strip() == text
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one heading {text!r}; found {len(matches)}")
    return matches[0]


def insert_after(document: Document, anchor, text: str, style: str | None = None):
    paragraph = document.add_paragraph(text, style=style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def intro_text() -> list[str]:
    raw = INTRO.read_text(encoding="utf-8")
    marker = "## The World of Diyse"
    if marker not in raw:
        raise RuntimeError("World intro reader section is missing")
    body = raw.split(marker, 1)[1]
    out = []
    for line in body.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("#"):
            break
        out.append(line.replace("**", "").replace("*", "").replace(chr(96), ""))
    if not out or out[-1] != "Mostly not metaphorically.":
        raise RuntimeError("World intro no longer ends on Nimera's approved sign-off")
    return out


def validate_authority() -> None:
    placement = PLACEMENT.read_text(encoding="utf-8")

    for authority, needles in FORMATION_AUTHORITIES:
        authority_text = authority.read_text(encoding="utf-8")
        for needle in needles:
            if needle not in authority_text:
                raise RuntimeError(
                    f"Current formation authority {authority.relative_to(ROOT)} no longer contains {needle!r}"
                )

    if "NO PLAYABLE MANDATORY CHAPTER-3 ROAD STRETCH" not in placement:
        raise RuntimeError("Chapter-3 no-road-combat firewall is missing")
    if "Redwater Initiate" not in placement or "Do not place it automatically" not in placement:
        raise RuntimeError("Waterworks Redwater placement firewall is missing")
    if "HOLD THE JUNCTION IS RETIRED" not in placement:
        raise RuntimeError("Hold-the-Junction retirement firewall is missing")

    cache = {}
    for _chapter, title, body, owner, needles in ROSTER_SUMMARIES:
        owner_text = cache.setdefault(owner, owner.read_text(encoding="utf-8"))
        for needle in needles:
            if needle not in owner_text:
                raise RuntimeError(f"{title}: current owning authority no longer contains {needle!r}")

    for _target, title, body, owner, needles in BRIDGES:
        owner_text = cache.setdefault(owner, owner.read_text(encoding="utf-8"))
        for needle in needles:
            if needle not in owner_text:
                raise RuntimeError(f"{title}: current owning authority no longer contains {needle!r}")
        joined = " ".join(body)
        if title == "Random Encounters — Old Waterworks" and "Redwater Initiate" in joined:
            raise RuntimeError("Redwater Initiate must not be auto-placed in Waterworks")
        if title.startswith("Random Encounters —") and "Way-Fort" in joined:
            raise RuntimeError("Way-Fort enemies must not be placed on the current mandatory Chapter-3 route")


def main() -> int:
    validate_authority()
    document = Document(READER)
    before = [p.text for p in document.paragraphs if is_dialogue(p)]
    if len(before) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(f"Reader has {len(before)} spoken lines; expected {EXPECTED_DIALOGUE_LINES}")

    for paragraph in document.paragraphs[:12]:
        text = paragraph.text.strip()
        if text == "Spoiler-Free Exact-Dialogue Reader":
            paragraph.runs[0].text = "Spoiler-Free Story & Gameplay Read-Through — Exact Dialogue"
        elif text.startswith("Current atomic dialogue edition."):
            paragraph.runs[0].text = (
                "Current Chapters 0–3 read-through. Spoken lines are copied verbatim from the active "
                "atomic dialogue sources; current world and gameplay bridges are restored from their "
                "owning authorities; writer-facing production notes are omitted."
            )

    for chapter_heading, title, body, _owner, _needles in ROSTER_SUMMARIES:
        anchor = find_heading(document, chapter_heading)
        anchor = insert_after(document, anchor, title, style="Heading 2")
        anchor.paragraph_format.space_before = Pt(6)
        anchor.paragraph_format.space_after = Pt(4)
        for text in body:
            anchor = insert_after(document, anchor, text)
            anchor.paragraph_format.space_after = Pt(3)

    chapter0 = find_heading(document, "Chapter 0")
    heading = chapter0.insert_paragraph_before("The World of Diyse", style="Heading 1")
    heading.paragraph_format.space_after = Pt(8)
    for text in intro_text():
        p = chapter0.insert_paragraph_before(text)
        p.paragraph_format.space_after = Pt(7)
    chapter0.paragraph_format.page_break_before = True

    for target_text, title, body, _owner, _needles in BRIDGES:
        target = find_heading(document, target_text)
        h = target.insert_paragraph_before(title, style="Heading 3")
        h.paragraph_format.space_before = Pt(7)
        h.paragraph_format.space_after = Pt(3)
        for text in body:
            p = target.insert_paragraph_before(text)
            p.paragraph_format.space_after = Pt(3)

    after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if after != before or len(after) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError("Reader enrichment altered spoken dialogue")

    document.save(READER)
    print(
        f"Reader enrichment complete: world intro + {len(ROSTER_SUMMARIES)} roster summaries + {len(BRIDGES)} encounter bridges; "
        f"{len(after)} spoken lines preserved exactly."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Restore reader-only layout invariants after presentation cleanup.

Cleanup removes production-facing paragraphs. Some source paragraphs can carry Word
section properties, so deleting them may collapse section breaks or detach the one shared
footer. This pass restores the intended reader presentation without touching spoken text:
Chapter 1–3 begin on fresh pages and the existing single footer part is reattached once.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
READER = ROOT / "build/dialogue/DIYSE_Chapters_0-3_Spoiler_Free_Exact_Dialogue_Reader_CURRENT.docx"
EXPECTED_DIALOGUE_LINES = 2165
LABEL_RE = re.compile(r"^.+:\s*$")
CHAPTER_BREAKS = {"Chapter 1", "Chapter 2", "Chapter 3"}


def is_dialogue(paragraph) -> bool:
    if not paragraph.runs:
        return False
    first = paragraph.runs[0]
    if not first.bold or not LABEL_RE.match(first.text):
        return False
    label = first.text.strip()[:-1].strip()
    return bool(re.search(r"[A-Z]", label)) and label == label.upper()


def main() -> int:
    if not READER.exists():
        raise RuntimeError(f"Generated reader does not exist: {READER.relative_to(ROOT)}")

    document = Document(READER)
    dialogue_before = [p.text for p in document.paragraphs if is_dialogue(p)]
    if len(dialogue_before) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError(
            f"Reader has {len(dialogue_before)} spoken lines; expected {EXPECTED_DIALOGUE_LINES}."
        )

    chapter_breaks_found: set[str] = set()
    for paragraph in document.paragraphs:
        if (
            paragraph.style
            and paragraph.style.name == "Heading 1"
            and paragraph.text.strip() in CHAPTER_BREAKS
        ):
            paragraph.paragraph_format.page_break_before = True
            chapter_breaks_found.add(paragraph.text.strip())

    if chapter_breaks_found != CHAPTER_BREAKS:
        missing = sorted(CHAPTER_BREAKS - chapter_breaks_found)
        raise RuntimeError(f"Reader chapter headings missing for page-break repair: {missing}")

    footer_rids = [
        rid for rid, relationship in document.part.rels.items()
        if relationship.reltype.endswith("/footer")
    ]
    if len(footer_rids) != 1:
        raise RuntimeError(
            f"Reader must contain exactly one shared footer relationship; found {footer_rids}."
        )

    # Cleanup may remove the paragraph that carried the first section's footerReference.
    # Reattach the already-generated footer part instead of creating another footer.
    for section in document.sections:
        sect_pr = section._sectPr
        for reference in list(sect_pr.findall(qn("w:footerReference"))):
            sect_pr.remove(reference)

    footer_reference = OxmlElement("w:footerReference")
    footer_reference.set(qn("w:type"), "default")
    footer_reference.set(qn("r:id"), footer_rids[0])
    document.sections[0]._sectPr.insert(0, footer_reference)

    dialogue_after = [p.text for p in document.paragraphs if is_dialogue(p)]
    if dialogue_before != dialogue_after or len(dialogue_after) != EXPECTED_DIALOGUE_LINES:
        raise RuntimeError("Reader layout repair altered spoken dialogue.")

    document.save(READER)

    # Package-level guard: exactly one footer part and exactly one explicit footer reference.
    with zipfile.ZipFile(READER) as archive:
        footer_parts = [
            name for name in archive.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        ]
        if len(footer_parts) != 1:
            raise RuntimeError(f"Expected one shared footer part; found {footer_parts}.")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_refs = re.findall(
            r'<w:footerReference[^>]+r:id="([^"]+)"',
            document_xml,
        )
        if footer_refs != [footer_rids[0]]:
            raise RuntimeError(
                f"Expected one explicit shared footer reference {footer_rids[0]}; found {footer_refs}."
            )

    print(
        "Reader layout repair complete; Chapter 1–3 start on fresh pages, "
        f"one shared footer is attached, and {len(dialogue_after)} spoken lines remain exact."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
